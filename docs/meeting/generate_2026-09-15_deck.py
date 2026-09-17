"""Generate the editable meeting deck from evidence-backed copy.

Run from the repository root:
    .venv/bin/python docs/meeting/generate_2026-09-15_deck.py
"""

from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt


HERE = Path(__file__).resolve().parent
OUTPUT = HERE / "2026-09-15-layered-la-agent-framework.pptx"

NAVY = RGBColor(15, 30, 48)
NAVY2 = RGBColor(25, 44, 64)
WHITE = RGBColor(248, 250, 250)
MUTED = RGBColor(173, 191, 201)
TEAL = RGBColor(63, 215, 192)
GOLD = RGBColor(246, 194, 93)
RED = RGBColor(245, 125, 115)
PANEL = RGBColor(28, 50, 70)
PANEL2 = RGBColor(35, 58, 77)
FONT = "Noto Sans CJK SC"

prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)


def shape(slide, x, y, w, h, fill=PANEL, line=None, radius=True):
    kind = MSO_SHAPE.ROUNDED_RECTANGLE if radius else MSO_SHAPE.RECTANGLE
    obj = slide.shapes.add_shape(kind, Inches(x), Inches(y), Inches(w), Inches(h))
    obj.fill.solid()
    obj.fill.fore_color.rgb = fill
    obj.line.fill.background() if line is None else None
    if line is not None:
        obj.line.color.rgb = line
    if radius:
        try:
            obj.adjustments[0] = 0.08
        except (IndexError, ValueError):
            pass
    return obj


def txt(slide, text, x, y, w, h, size=20, color=WHITE, bold=False,
        align=PP_ALIGN.LEFT, valign=MSO_ANCHOR.MIDDLE, margin=0.02):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame
    tf.clear()
    tf.word_wrap = True
    tf.margin_left = Inches(margin)
    tf.margin_right = Inches(margin)
    tf.margin_top = Inches(0.015)
    tf.margin_bottom = Inches(0.015)
    tf.vertical_anchor = valign
    for i, line in enumerate(text.split("\n")):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.text = line
        p.alignment = align
        p.space_after = Pt(2)
        p.font.name = FONT
        p.font.size = Pt(size)
        p.font.bold = bold
        p.font.color.rgb = color
    return box


def line(slide, x1, y1, x2, y2, color=MUTED, width=1.4):
    sh = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(x1), Inches(y1),
                                    Inches(x2), Inches(y2))
    sh.line.color.rgb = color
    sh.line.width = Pt(width)
    return sh


def base(title, kicker, source, number):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = NAVY
    shape(slide, 0, 0, 0.14, 7.5, TEAL, radius=False)
    txt(slide, kicker.upper(), 0.5, 0.24, 10.5, 0.26, 10, TEAL, True)
    txt(slide, title, 0.5, 0.63, 12.15, 0.58, 26, WHITE, True)
    line(slide, 0.5, 1.33, 12.83, 1.33, RGBColor(58, 79, 96), 1)
    txt(slide, source, 0.5, 7.12, 11.8, 0.21, 8, MUTED)
    txt(slide, f"{number:02d}", 12.25, 7.07, 0.5, 0.28, 10, TEAL, True,
        PP_ALIGN.RIGHT)
    return slide


def card(slide, x, y, w, h, title, body, accent=TEAL, body_size=16,
         title_size=18):
    shape(slide, x, y, w, h, PANEL)
    shape(slide, x + 0.15, y + 0.18, 0.06, 0.38, accent, radius=False)
    txt(slide, title, x + 0.31, y + 0.15, w - 0.48, 0.46,
        title_size, WHITE, True)
    txt(slide, body, x + 0.22, y + 0.73, w - 0.44, h - 0.89,
        body_size, MUTED, valign=MSO_ANCHOR.TOP)


def pill(slide, label, x, y, w, fill=PANEL2, color=WHITE, size=12):
    shape(slide, x, y, w, 0.42, fill)
    txt(slide, label, x + 0.05, y + 0.02, w - 0.1, 0.36, size, color, True,
        PP_ALIGN.CENTER)


# 1 — title
s = prs.slides.add_slide(prs.slide_layouts[6])
s.background.fill.solid()
s.background.fill.fore_color.rgb = NAVY
shape(s, 0, 0, 0.16, 7.5, TEAL, radius=False)
shape(s, 8.4, 0.55, 4.5, 6.35, NAVY2)
txt(s, "分层 LA", 0.62, 0.8, 7.0, 0.64, 22, TEAL, True)
txt(s, "具身 Agent 框架", 0.62, 1.58, 7.15, 0.9, 39, WHITE, True)
txt(s, "从开放目标到可验证的机器人任务", 0.62, 2.66, 7.12, 0.55,
    21, MUTED)
shape(s, 0.62, 3.52, 6.9, 0.085, TEAL, radius=False)
txt(s, "架构实现 · 首版任务拆分 · 阶段 1 指标 · 我的工作", 0.62, 3.87,
    7.1, 0.75, 18, WHITE)
txt(s, "会议讨论材料  |  2026.09.15", 0.62, 6.6, 7.0, 0.32,
    12, MUTED)
for y, n, title, sub in [
    (1.08, "01", "目标与规划", "秒—分钟：解释目标、生成 DAG"),
    (2.66, "02", "确定性执行", "毫秒—秒：契约、资源、恢复"),
    (4.24, "03", "控制与证据", "实时控制、独立验证与安全"),
]:
    pill(s, n, 8.82, y, 0.62, TEAL, NAVY, 13)
    txt(s, title, 9.63, y - 0.02, 2.95, 0.43, 21, WHITE, True)
    txt(s, sub, 9.63, y + 0.49, 2.95, 0.45, 13, MUTED)


# 2 — thesis
s = base("一句话定位：模型提议，框架裁决，机器人执行", "01 / 定位",
         "依据：docs/DETAILED_ARCHITECTURE.md；docs/ARCHITECTURE.md", 2)
card(s, 0.55, 1.75, 3.9, 3.82, "上层：LA / 语义与计划",
     "解释用户目标，构造上下文，生成候选任务 DAG。\n\n不直接发关节/速度命令。", TEAL, 18)
card(s, 4.72, 1.75, 3.9, 3.82, "中层：确定性 Runtime",
     "校验计划，冻结技能，管理依赖、资源、执行 ID、恢复与结果。", GOLD, 18)
card(s, 8.89, 1.75, 3.9, 3.82, "底层：控制器 + 安全",
     "ROS2/机器人适配器执行动作。阶段 1 目标：控制端拒绝旧权威；硬件安全链路独立。", RED, 18)
shape(s, 0.55, 5.87, 12.24, 0.76, NAVY2)
txt(s, "判定原则：HTTP 200 ≠ 技能成功；模型计划合法 ≠ 机器人目标完成。",
    0.82, 6.04, 11.75, 0.42, 19, WHITE, True)


# 3 — planes
s = base("分层目标架构：四个平面，各有权威数据", "02 / 总体架构",
         "依据：docs/DETAILED_ARCHITECTURE.md §3；docs/PROJECT_STATUS.md", 3)
rows = [
    ("Agent 控制平面", "Goal → Context → Planner → Validator → Runtime", "Task / Plan / Execution / Resource", TEAL),
    ("机器人适配平面", "Skill Registry → HTTP / ROS2 Action → 控制器", "远端执行状态 / 动作反馈", GOLD),
    ("证据数据平面", "事件账本 → MCAP / rosbag2 → Asset / Memory → UI", "观测 / 资产 / 验证来源", RGBColor(135, 178, 241)),
    ("独立安全平面", "限幅 / 急停 / Watchdog / 控制权校验", "物理安全状态；不依赖 LLM", RED),
]
for i, (title, chain, owner, accent) in enumerate(rows):
    y = 1.67 + i * 1.25
    shape(s, 0.6, y, 12.14, 1.01, PANEL)
    shape(s, 0.6, y, 0.10, 1.01, accent, radius=False)
    txt(s, title, 0.9, y + 0.16, 2.43, 0.43, 18, WHITE, True)
    txt(s, chain, 3.12, y + 0.08, 6.32, 0.79, 16, WHITE)
    txt(s, owner, 9.5, y + 0.12, 2.96, 0.71, 13, MUTED)


# 4 — status
s = base("当前站位：软件骨架可运行，机器人闭环尚未验收", "03 / 现状",
         "依据：docs/PROJECT_STATUS.md（2026-09-12）；docs/VALIDATION.md", 4)
card(s, 0.58, 1.67, 5.88, 4.74, "已有实现与证据",
     "• DAG 契约、计划历史与有界重规划\n• 单机多任务、多资源原子分配\n• DBOS 恢复 + 固定 execution ID\n• 本地文件与 HTTP 技能协议\n• 模型响应、资产、事件和只读 UI",
     TEAL, 18)
card(s, 6.75, 1.67, 5.88, 4.74, "关键缺口",
     "• 控制端 fencing、补偿和人工接管\n• 真实 ROS2 / Unitree 技能及停止回执\n• 独立目标 verifier 的真机证据\n• 真实传感器与地图失效链路\n• 长程规划质量、多机器人/生产 HA",
     GOLD, 18)


# 5 — planning pipeline
s = base("首版任务拆分：受约束的 DAG 生成 + 双层校验", "04 / 核心机制",
         "依据：robot_agent/runtime.py、planner.py、model_methods.py、contracts.py", 5)
steps = [
    ("01", "输入", "原始 goal\n注册技能目录"),
    ("02", "上下文", "GoalContext\n能力 / 任务状态"),
    ("03", "模型候选", "JSON steps\ndeps + verification"),
    ("04", "计划门禁", "DAG / 引用\n最终 verifier"),
    ("05", "执行门禁", "资源 / 重入\n冻结 plan + catalog"),
]
for i, (no, title, body) in enumerate(steps):
    x = 0.58 + i * 2.51
    shape(s, x, 2.07, 2.24, 2.85, PANEL)
    pill(s, no, x + 0.16, 2.26, 0.55, TEAL, NAVY, 11)
    txt(s, title, x + 0.16, 2.84, 1.88, 0.5, 19, WHITE, True)
    txt(s, body, x + 0.16, 3.54, 1.91, 1.04, 15, MUTED)
    if i < 4:
        txt(s, "→", x + 2.25, 3.16, 0.26, 0.4, 22, TEAL, True,
            PP_ALIGN.CENTER)
shape(s, 0.58, 5.33, 12.28, 0.92, NAVY2)
txt(s, "当前常规 agent 路径直接用原始 goal 规划；独立语义分析入口尚未自动接入提交链。",
    0.82, 5.58, 11.82, 0.39, 16, GOLD, True)


# 6 — DAG example
s = base("一条真实模型计划：两步，但尚未执行", "05 / 可追溯样例",
         "依据：docs/validation/model-planning-qwen3.8-max.json（2026-09-07）", 6)
shape(s, 0.67, 1.85, 4.55, 3.41, PANEL)
pill(s, "STEP 1", 0.98, 2.13, 1.06, TEAL, NAVY)
txt(s, "ingest_readme", 0.97, 2.81, 3.87, 0.57, 23, WHITE, True)
txt(s, "file.ingest\n计划动作：读取 README，返回 asset_id",
    0.97, 3.57, 3.88, 1.12, 17, MUTED)
txt(s, "asset_id →", 5.32, 3.27, 2.55, 0.53, 22, GOLD, True,
    PP_ALIGN.CENTER)
shape(s, 8.1, 1.85, 4.55, 3.41, PANEL)
pill(s, "VERIFIER", 8.42, 2.13, 1.42, GOLD, NAVY)
txt(s, "verify_readme_asset", 8.41, 2.81, 3.88, 0.57, 21, WHITE, True)
txt(s, "asset.verify\n校验资产并作为最终验证节点",
    8.41, 3.57, 3.89, 1.12, 17, MUTED)
shape(s, 0.68, 5.62, 11.97, 0.77, NAVY2)
txt(s, "实测结论：模型返回并通过计划契约；plan_executed = false。",
    0.97, 5.78, 11.25, 0.43, 18, WHITE, True)


# 7 — validation
s = base("计划不是字符串：必须经过确定性门禁", "06 / 校验规则",
         "依据：robot_agent/contracts.py、planner.py、runtime.py", 7)
card(s, 0.62, 1.72, 3.92, 4.54, "结构合法性",
     "1–256 步、ID 唯一\n依赖存在且无环\n所有步骤通向最终验证\n引用仅指向直接依赖的输出",
     TEAL, 18)
card(s, 4.71, 1.72, 3.92, 4.54, "能力与权限",
     "技能必须已注册\n模型只见允许暴露的能力字段\n最终节点必须是 verifier\n重试仅用于安全可重入技能",
     GOLD, 18)
card(s, 8.8, 1.72, 3.92, 4.54, "接受与审计",
     "当前资源容量再检查\n冻结计划和技能目录\n记录原始模型响应\n无效候选拒绝，不造默认计划",
     RGBColor(135, 178, 241), 18)


# 8 — execution
s = base("任务执行闭环：身份、资源和真实证据", "07 / Runtime",
         "依据：docs/ARCHITECTURE.md；tests/test_http_execution.py、test_durable_planner.py", 8)
for i, (head, body, accent) in enumerate([
    ("派发", "依赖成功 + 资源可得；外部 IO 前保存 execution ID", TEAL),
    ("执行", "PUT execute / GET query / POST cancel；按原 ID 恢复", GOLD),
    ("判定", "终态 + quiescent + schema + 后置条件 + 来源证据", RGBColor(135, 178, 241)),
    ("异常", "unknown 不重发、不释放资源；取消优先于迟到重规划", RED),
]):
    y = 1.65 + i * 1.22
    shape(s, 0.7, y, 11.9, 1.02, PANEL)
    pill(s, f"0{i+1}", 0.94, y + 0.28, 0.59, accent, NAVY, 11)
    txt(s, head, 1.77, y + 0.18, 1.28, 0.59, 19, WHITE, True)
    txt(s, body, 3.11, y + 0.17, 9.05, 0.62, 17, MUTED)


# 9 — stage 1
s = base("路线图阶段 1：执行权威与恢复闭环", "08 / 当前阶段",
         "依据：docs/BUILD_ROADMAP.md §阶段 1；docs/PROJECT_STATUS.md", 9)
items = [
    ("已完成", "历史计划快照", "plan / steps / catalog / revision / generation", TEAL),
    ("下一项", "控制端 fencing", "单调 token；旧执行者不能再推进副作用", GOLD),
    ("随后", "补偿与人工接管", "补偿条件、实际执行、结果与人工终态", RGBColor(135, 178, 241)),
    ("随后", "重试归属 + 故障矩阵", "防重复副作用；覆盖断进程、超时、迟到响应等", RED),
]
for i, (status, title, body, accent) in enumerate(items):
    y = 1.62 + i * 1.23
    shape(s, 0.68, y, 11.98, 1.05, PANEL)
    pill(s, status, 0.94, y + 0.28, 1.05, accent, NAVY, 12)
    txt(s, title, 2.2, y + 0.17, 3.1, 0.57, 20, WHITE, True)
    txt(s, body, 5.4, y + 0.16, 6.84, 0.64, 16, MUTED)


# 10 — metrics
s = base("把“已有证据”和“建议指标”放在两列", "09 / 验收口径",
         "依据：2026-09-14 本地 pytest；docs/VALIDATION.md；指标为会议建议值", 10)
headers = [("维度", 0.7, 2.16), ("当前证据", 2.94, 4.19),
           ("建议下一轮验收", 7.2, 5.42)]
for h, x, w in headers:
    txt(s, h, x + 0.14, 1.69, w - 0.2, 0.42, 15, TEAL, True)
data = [
    ("软件回归", "63 passed / 1 skipped / 9 subtests", "全量通过 + 故障用例可重现"),
    ("模型拆分", "1 个真实两步合法 DAG，未执行", "3 类×10 条目标；分别统计合法/可执行/验证率"),
    ("恢复/资源", "同一 ID 恢复；unknown 不释放", "测试矩阵中重复副作用 0、资源早释 0"),
    ("控制权", "账本 generation 已有", "旧 token 副作用接受数 0"),
    ("真机完成", "暂无机器人端到端证据", "阶段 2：至少 1 项技能的完整纵向切片"),
]
for i, row in enumerate(data):
    y = 2.21 + i * 0.82
    shape(s, 0.68, y, 11.97, 0.72, PANEL if i % 2 == 0 else PANEL2)
    txt(s, row[0], 0.82, y + 0.08, 1.95, 0.52, 15, WHITE, True)
    txt(s, row[1], 3.04, y + 0.08, 3.98, 0.52, 14, MUTED)
    txt(s, row[2], 7.29, y + 0.07, 5.08, 0.55, 14, WHITE)
txt(s, "建议数值只定义测试集/故障矩阵的通过门槛，不代表现场可靠性已测得。",
    0.72, 6.55, 11.74, 0.35, 13, GOLD)


# 11 — personal work
s = base("我的建议主责：计划契约、执行权威、证据链", "10 / 个人工作",
         "本页为会议建议分工，需与机器人控制、安全和场景负责人确认", 11)
card(s, 0.62, 1.63, 5.92, 2.18, "P0｜任务拆分契约与评测",
     "固化 Goal→Context→Plan；建 3 类×10 条目标集；记录每次原始响应、拒绝原因和分项统计。",
     TEAL, 16, 17)
card(s, 6.78, 1.63, 5.92, 2.18, "P0｜阶段 1 控制权",
     "实现 token/执行记录/协议校验；覆盖重启、取消竞争、旧执行者和资源早释故障矩阵。",
     GOLD, 16, 17)
card(s, 0.62, 4.03, 5.92, 2.18, "P1｜真实技能纵向切片",
     "与机器人同事定义 1 项技能的 execute/query/cancel、schema、资源、停止回执和 verifier。",
     RGBColor(135, 178, 241), 16, 17)
card(s, 6.78, 4.03, 5.92, 2.18, "P1｜证据与演示",
     "维护状态和验证记录；演示从目标到计划、execution ID、实际结果和证据来源的追溯。",
     RED, 16, 17)


# 12 — ask
s = base("这次会议希望确认的三件事", "11 / 讨论与决策",
         "依据：docs/OPEN_QUESTIONS.md；docs/BUILD_ROADMAP.md", 12)
card(s, 0.65, 1.76, 3.87, 3.73, "01｜第一批机器人技能",
     "选哪一个任务作为低风险纵向切片？谁提供 ROS2/Unitree 执行、查询、取消和独立观测？",
     TEAL, 19)
card(s, 4.73, 1.76, 3.87, 3.73, "02｜控制端接口",
     "控制器能否校验 fencing token 并返回真实停止回执？失败和 unknown 由谁确认？",
     GOLD, 19)
card(s, 8.81, 1.76, 3.87, 3.73, "03｜任务评测口径",
     "先评哪 3 类目标？合法计划、可执行计划、最终完成分别以什么证据为准？",
     RGBColor(135, 178, 241), 19)
shape(s, 0.65, 5.78, 12.03, 0.7, NAVY2)
txt(s, "下一步：先关住旧执行者，再扩大自主规划与真机任务范围。",
    0.91, 5.91, 11.44, 0.43, 19, WHITE, True)


# 13 — appendix
s = base("证据索引：现场追问时打开这些文件", "附录 / 证据",
         "材料依据为当前仓库文件；未新增真实模型或机器人实验", 13)
sources = [
    ("架构与边界", "docs/DETAILED_ARCHITECTURE.md · docs/ARCHITECTURE.md"),
    ("实际完成状态", "docs/PROJECT_STATUS.md · docs/BUILD_ROADMAP.md"),
    ("拆分机制", "docs/TASK_DECOMPOSITION_FLOW.md · robot_agent/planner.py"),
    ("模型证据", "docs/validation/model-planning-qwen3.8-max.json"),
    ("运行验收", "docs/VALIDATION.md · tests/test_http_execution.py"),
    ("语义上下文", "tests/test_semantic_context_pipeline.py"),
]
for i, (head, body) in enumerate(sources):
    y = 1.62 + i * 0.77
    shape(s, 0.66, y, 12.0, 0.63, PANEL if i % 2 == 0 else PANEL2)
    txt(s, head, 0.91, y + 0.09, 2.18, 0.41, 15, WHITE, True)
    txt(s, body, 3.05, y + 0.09, 9.15, 0.41, 14, MUTED)


prs.save(OUTPUT)
print(OUTPUT)
